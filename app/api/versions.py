import io
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlmodel import Session, select
from docx import Document

from app.db import get_session
from app.models import ChapterDraft, ChapterReview, EditorAction, LLMCall, OutlineRevision, Project, SourceChapter, StoryBibleSnapshot, Version
from app.schemas import ActionInput, EditOutline, EditChapter, Instruction, Character, Direction, EditDirection
from app.checks import chapter_count, outline_issues, validate_new_characters, clean_draft, count_words, validate_direction
from app.services import rows, latest_outline, latest_plans, current_drafts, initial_bible, version_bible, log_action
from app.pipeline.runner import runner

router = APIRouter()


def require_version(ident, session):
    version = session.get(Version, ident)
    if not version:
        raise HTTPException(404, "找不到续写版本")
    return version


def editable(version, statuses):
    if runner.busy(version.project_id, version.id) or version.status not in statuses:
        raise HTTPException(409, "当前阶段不能执行该操作")


@router.get("/versions/{ident}")
def detail(ident: int, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    history = rows(ChapterDraft, ChapterDraft.version_id == ident, order=ChapterDraft.index)
    ids = [d.id for d in history]
    return {"version": version, "outline": latest_outline(ident), "outline_history": rows(OutlineRevision, OutlineRevision.version_id == ident, order=OutlineRevision.rev), "plans": latest_plans(ident), "chapters": current_drafts(ident), "draft_history": history, "reviews": rows(ChapterReview, ChapterReview.draft_id.in_(ids)), "bible": version_bible(ident).model_dump(), "actions": rows(EditorAction, EditorAction.version_id == ident, order=EditorAction.id), "calls": rows(LLMCall, LLMCall.version_id == ident, order=LLMCall.id)}


@router.put("/versions/{ident}/direction")
async def edit_direction(ident: int, body: EditDirection, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"awaiting_direction_confirmation"})
    try:
        validate_direction(body.direction, initial_bible(version.project_id))
        if not body.direction.opening_moment.strip():
            raise ValueError("请填写从原文结尾哪个瞬间接起")
    except ValueError as error:
        raise HTTPException(422, str(error))
    before = version.direction
    version.direction = body.direction.model_dump()
    session.add(version)
    log_action(session, version.project_id, ident, "direction_edit", body.editor, before=before, after=version.direction)
    session.commit()
    return version


@router.post("/versions/{ident}/direction/confirm", status_code=202)
async def confirm_direction(ident: int, body: ActionInput, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"awaiting_direction_confirmation"})
    version.status, version.current_node = "queued", "N4"
    session.add(version)
    log_action(session, version.project_id, ident, "direction_confirm", body.editor, direction=version.direction)
    session.commit()
    runner.enqueue(version.project_id, ident)
    return {"status": "queued"}


@router.put("/versions/{ident}/outline")
def edit_outline(ident: int, body: EditOutline, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"awaiting_outline"})
    project = session.get(Project, version.project_id)
    sources = rows(SourceChapter, SourceChapter.project_id == project.id, order=SourceChapter.index)
    count = chapter_count(version.target_words, project.channel)
    if [c.index for c in body.outline.chapters] != list(range(sources[-1].index + 1, sources[-1].index + count + 1)):
        raise HTTPException(422, "请保持指定章节数与连续章号")
    try:
        validate_new_characters(body.outline.new_characters, initial_bible(project.id))
    except ValueError as error:
        raise HTTPException(422, str(error))
    old = latest_outline(ident)
    revision = OutlineRevision(version_id=ident, rev=old.rev + 1, data=body.outline.model_dump(), author=body.editor, warnings=outline_issues(body.outline, Direction.model_validate(version.direction), sources[-1].index + 1, count))
    session.add(revision)
    log_action(session, project.id, ident, "outline_edit", body.editor, before=old.data, after=revision.data)
    session.commit()
    session.refresh(revision)
    return revision


@router.post("/versions/{ident}/outline/confirm", status_code=202)
async def confirm_outline(ident: int, body: ActionInput, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"awaiting_outline"})
    outline = latest_outline(ident)
    bible = initial_bible(version.project_id)
    for item in outline.data["new_characters"]:
        bible.characters.append(Character(name=item["name"], role="supporting", goal_now=item["purpose"]))
    notes = {op["id"]: f"第{c['index']}章：{op['note']}" for c in outline.data["chapters"] for op in c["foreshadow_ops"] if op["op"] == "payoff"}
    for foreshadow in bible.foreshadows:
        if foreshadow.id in notes:
            foreshadow.plan_note = notes[foreshadow.id]
    start = max(c.index for c in rows(SourceChapter, SourceChapter.project_id == version.project_id))
    session.add(StoryBibleSnapshot(project_id=version.project_id, version_id=ident, after_chapter=start, data=bible.model_dump()))
    outline.confirmed = True
    version.status, version.current_node = "queued", "N5"
    session.add_all([outline, version])
    log_action(session, version.project_id, ident, "outline_confirm", body.editor, revision=outline.rev, warnings=outline.warnings, new_characters=outline.data["new_characters"])
    session.commit()
    runner.enqueue(version.project_id, ident)
    return {"status": "queued"}


def find_draft(ident, index):
    draft = next((d for d in current_drafts(ident) if d.index == index), None)
    if not draft:
        raise HTTPException(404, "找不到该章节")
    return draft


def mark_edit(version, index, session):
    version.stale_chapters = sorted(set(version.stale_chapters) | {d.index for d in current_drafts(version.id) if d.index > index})
    version.status, version.current_node = "queued", "N6_EDIT"
    session.add(version)


@router.put("/versions/{ident}/chapters/{index}")
async def edit_chapter(ident: int, index: int, body: EditChapter, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"awaiting_review"})
    old = find_draft(ident, index)
    all_drafts = rows(ChapterDraft, ChapterDraft.version_id == ident, ChapterDraft.index == index)
    old.is_current = False
    content = clean_draft(body.content, index, old.title)
    draft = ChapterDraft(version_id=ident, index=index, rev=max(d.rev for d in all_drafts) + 1, title=old.title, content=content, word_count=count_words(content.partition("\n")[2]), author=body.editor)
    session.add_all([old, draft])
    mark_edit(version, index, session)
    log_action(session, version.project_id, ident, "chapter_edit", body.editor, index=index, before=old.content, after=content)
    session.commit()
    session.refresh(draft)
    runner.enqueue(version.project_id, ident)
    return draft


@router.post("/versions/{ident}/chapters/{index}/regenerate", status_code=202)
async def regenerate(ident: int, index: int, body: Instruction, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"awaiting_review"})
    old = find_draft(ident, index)
    mark_edit(version, index, session)
    log_action(session, version.project_id, ident, "chapter_regenerate", body.editor, index=index, instruction=body.instruction, before_draft_id=old.id)
    session.commit()
    runner.enqueue(version.project_id, ident)
    return {"status": "queued"}


@router.post("/versions/{ident}/chapters/{index}/accept")
def accept(ident: int, index: int, body: ActionInput, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"awaiting_review"})
    draft = find_draft(ident, index)
    reviews = rows(ChapterReview, ChapterReview.draft_id == draft.id, order=ChapterReview.id.desc())
    if not reviews:
        raise HTTPException(409, "审核尚未完成")
    review = reviews[0]
    review.editor_accepted = True
    version.stale_chapters = [n for n in version.stale_chapters if n != index]
    session.add_all([review, version])
    log_action(session, version.project_id, ident, "chapter_accept", body.editor, index=index, draft_id=draft.id, verdict=review.verdict, reason=body.reason)
    session.commit()
    return {"status": "accepted"}


@router.post("/versions/{ident}/retry", status_code=202)
async def retry(ident: int, body: ActionInput, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"failed"})
    version.status, version.error = "queued", None
    session.add(version)
    log_action(session, version.project_id, ident, "version_retry", body.editor, node=version.current_node)
    session.commit()
    runner.enqueue(version.project_id, ident)
    return {"status": "queued"}


@router.post("/versions/{ident}/submit")
def submit(ident: int, body: ActionInput, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"awaiting_review"})
    drafts = current_drafts(ident)
    if len(drafts) != version.total_chapters:
        raise HTTPException(409, "还有章节未生成")
    if version.stale_chapters:
        raise HTTPException(409, "请先核对被标记为可能不连贯的章节")
    for draft in drafts:
        reviews = rows(ChapterReview, ChapterReview.draft_id == draft.id, order=ChapterReview.id.desc())
        if not reviews or (reviews[0].verdict != "pass" and not reviews[0].editor_accepted):
            raise HTTPException(409, f"第{draft.index}章存在审核意见，请修改或人工确认")
    version.status = "submitted"
    session.add(version)
    log_action(session, version.project_id, ident, "version_submit", body.editor, draft_ids=[d.id for d in drafts])
    session.commit()
    return {"status": "submitted"}


@router.post("/versions/{ident}/abandon")
def abandon(ident: int, body: ActionInput, session: Session = Depends(get_session)):
    version = require_version(ident, session)
    editable(version, {"failed", "awaiting_outline", "awaiting_review", "awaiting_direction_confirmation"})
    version.status = "abandoned"
    session.add(version)
    log_action(session, version.project_id, ident, "version_abandon", body.editor, reason=body.reason)
    session.commit()
    return {"status": "abandoned"}


@router.get("/versions/{ident}/export")
def export(ident: int, format: str = "txt", session: Session = Depends(get_session)):
    version = require_version(ident, session)
    if format not in {"txt", "docx"}:
        raise HTTPException(422, "只支持txt或docx")
    project = session.get(Project, version.project_id)
    sources = rows(SourceChapter, SourceChapter.project_id == project.id, order=SourceChapter.index)
    text = "\n\n".join([f"{c.title}\n\n{c.content}" for c in sources] + [d.content for d in current_drafts(ident)])
    if project.is_demo:
        text = "【演示数据，仅用于验证操作流程】\n\n" + text
    if format == "txt":
        content, media = text.encode("utf-8-sig"), "text/plain; charset=utf-8"
    else:
        doc = Document()
        doc.add_heading(project.title, 0)
        for line in text.splitlines():
            doc.add_paragraph(line)
        buf = io.BytesIO()
        doc.save(buf)
        content, media = buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return Response(content=content, media_type=media, headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(project.title, safe='')}-v{ident}.{format}"})
